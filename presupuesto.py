import customtkinter as ctk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from database import Database
from utilidades import formatear_fecha
from colores_app import *
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

MESES_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

class GestionPresupuestos:
    def __init__(self, padre, rol, usuario_actual):
        self.padre = padre
        self.rol = rol
        self.usuario_actual = usuario_actual
        self.bd = Database()
        self.usuario_id = self.bd.obtener_id_usuario(usuario_actual) or 0

        self.marco = ctk.CTkFrame(padre, fg_color=FONDO_TARJETA)
        self.marco.pack(fill="both", expand=True, padx=10, pady=10)

        self.tabview = ctk.CTkTabview(self.marco, fg_color=FONDO_SIDEBAR)
        self.tabview.pack(fill="both", expand=True, padx=5, pady=5)

        self.pestania_ordenes = self.tabview.add("📋 Cuentas por Cobrar")
        self.pestania_historial = self.tabview.add("📜 Historial de Pagos")
        self.pestania_graficas = self.tabview.add("📊 Estadísticas Financieras")

        self.configurar_pestania_ordenes()
        self.configurar_pestania_historial()
        self.configurar_pestania_graficas()

    def configurar_pestania_ordenes(self):
        barra_herramientas = ctk.CTkFrame(self.pestania_ordenes, fg_color="transparent")
        barra_herramientas.pack(fill="x", pady=5)

        boton_refrescar = ctk.CTkButton(
            barra_herramientas, text="⟳ Actualizar Lista",
            fg_color=COLOR_ACENTO, text_color=TEXTO_BLANCO,
            command=self.cargar_ordenes
        )
        boton_refrescar.pack(side="left", padx=5)

        etiqueta_instruccion = ctk.CTkLabel(
            barra_herramientas, text="💡 Haz doble clic en una orden para ver detalle",
            text_color=TEXTO_GRIS, font=("Inter", 11)
        )
        etiqueta_instruccion.pack(side="right", padx=10)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview", background=FONDO_TARJETA, foreground=TEXTO_BLANCO,
                        fieldbackground=FONDO_TARJETA, rowheight=25)
        style.map("Treeview", background=[('selected', COLOR_ACENTO)])

        self.arbol_ordenes = ttk.Treeview(
            self.pestania_ordenes,
            columns=("ID", "Cliente", "Vehículo", "Total USD", "Pagado USD", "Saldo USD", "Estado"),
            show="headings"
        )
        columnas = [
            ("ID", 50), ("Cliente", 150), ("Vehículo", 150),
            ("Total USD", 100), ("Pagado USD", 100), ("Saldo USD", 100), ("Estado", 100)
        ]
        for columna, width in columnas:
            self.arbol_ordenes.heading(columna, text=columna)
            self.arbol_ordenes.column(columna, width=width, anchor="center" if "USD" in columna or columna in ["ID", "Estado"] else "w")

        self.arbol_ordenes.tag_configure("pagado", foreground="#2ecc71")
        self.arbol_ordenes.tag_configure("pendiente", foreground="#e74c3c")

        scroll = ttk.Scrollbar(self.pestania_ordenes, orient="vertical", command=self.arbol_ordenes.yview)
        self.arbol_ordenes.configure(yscrollcommand=scroll.set)
        self.arbol_ordenes.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

        self.arbol_ordenes.bind("<Double-Button-1>", self.on_orden_seleccionada)
        self.cargar_ordenes()

    def cargar_ordenes(self):
        for row in self.arbol_ordenes.get_children():
            self.arbol_ordenes.delete(row)

        ordenes = self.bd.listar_cuentas_por_cobrar()
        if ordenes is None:
            messagebox.showerror(
                "Error de base de datos",
                "No se pudieron cargar las cuentas por cobrar.\nRevise la conexión con la base de datos.",
                parent=self.marco
            )
            return

        for o in ordenes:
            saldo = float(o['saldo'])
            estado = "Saldado" if saldo <= 0.01 else "Pendiente"
            tag = "pagado" if estado == "Saldado" else "pendiente"
            self.arbol_ordenes.insert("", "end", values=(
                o['id'],
                o['cliente'],
                o['vehiculo'],
                f"${float(o['total_orden_usd']):.2f}",
                f"${float(o['pagado']):.2f}",
                f"${max(0, saldo):.2f}",
                estado
            ), tags=(tag,))

    def configurar_pestania_historial(self):
        barra_herramientas = ctk.CTkFrame(self.pestania_historial, fg_color="transparent")
        barra_herramientas.pack(fill="x", pady=5)

        boton_refrescar = ctk.CTkButton(
            barra_herramientas, text="⟳ Refrescar Historial",
            fg_color=COLOR_AZUL, text_color=TEXTO_BLANCO,
            command=self.cargar_historial
        )
        boton_refrescar.pack(side="left", padx=5)

        boton_reporte = ctk.CTkButton(
            barra_herramientas, text="💵 Reporte de Recaudos",
            fg_color=COLOR_VERDE, text_color=TEXTO_BLANCO,
            command=self.abrir_reporte_recaudos
        )
        boton_reporte.pack(side="left", padx=5)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview", background=FONDO_TARJETA, foreground=TEXTO_BLANCO,
                        fieldbackground=FONDO_TARJETA, rowheight=25)
        style.map("Treeview", background=[('selected', COLOR_ACENTO)])

        self.arbol_historial = ttk.Treeview(
            self.pestania_historial,
            columns=("ID", "Orden", "Cliente", "Vehículo", "Monto Original", "Moneda", "Tasa", "Monto USD", "Fecha", "Método", "Referencia"),
            show="headings"
        )
        columnas = [
            ("ID", 40), ("Orden", 50), ("Cliente", 150), ("Vehículo", 120),
            ("Monto Original", 100), ("Moneda", 60), ("Tasa", 80),
            ("Monto USD", 100), ("Fecha", 120), ("Método", 100), ("Referencia", 100)
        ]
        for columna, width in columnas:
            self.arbol_historial.heading(columna, text=columna)
            self.arbol_historial.column(columna, width=width, anchor="center" if "USD" in columna or columna in ["ID", "Orden", "Monto Original", "Monto USD"] else "w")

        scroll = ttk.Scrollbar(self.pestania_historial, orient="vertical", command=self.arbol_historial.yview)
        self.arbol_historial.configure(yscrollcommand=scroll.set)
        self.arbol_historial.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

        self.cargar_historial()

    def cargar_historial(self):
        for row in self.arbol_historial.get_children():
            self.arbol_historial.delete(row)

        pagos = self.bd.obtener_historial_pagos()
        if pagos is None:
            messagebox.showerror(
                "Error de base de datos",
                "No se pudo cargar el historial de pagos.\nRevise la conexión con la base de datos.",
                parent=self.marco
            )
            return

        total_usd = 0
        for p in pagos:
            total_usd += p['monto_ref_usd']
            self.arbol_historial.insert("", "end", values=(
                p['id'],
                p['orden_id'],
                p['cliente'],
                p['vehiculo'],
                f"{p['monto_original']:.2f}",
                p['moneda'],
                f"{p['tasa_cambio']:.2f}",
                f"{p['monto_ref_usd']:.2f}",
                formatear_fecha(p["fecha_pago"]),
                p['metodo_pago'],
                p['referencia'] or ""
            ))

        self.etiqueta_total_historial = ctk.CTkLabel(
            self.pestania_historial,
            text=f"💰 Total recaudado: ${total_usd:.2f} USD",
            font=("Inter", 14, "bold"),
            text_color=COLOR_VERDE
        )
        self.etiqueta_total_historial.pack(anchor="e", padx=10, pady=5)

    def configurar_pestania_graficas(self):
        self.marco_graficas_container = ctk.CTkFrame(self.pestania_graficas, fg_color="transparent")
        self.marco_graficas_container.pack(fill="both", expand=True)
        self.renderizar_graficas()

    def renderizar_graficas(self):
        for widget in self.marco_graficas_container.winfo_children():
            widget.destroy()

        pagos_mensuales = self.bd.obtener_resumen_pagos()
        distribucion = self.bd.obtener_distribucion_monedas()

        fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(9, 4))
        fig.patch.set_facecolor(FONDO_TARJETA)

        ax1.set_facecolor(FONDO_TARJETA)
        if pagos_mensuales:
            meses = [p['mes'] for p in pagos_mensuales]
            totales = [float(p['total_usd']) for p in pagos_mensuales]
            ax1.bar(meses, totales, color=COLOR_ACENTO)
            ax1.set_title("Ingresos por Mes (USD)", color=TEXTO_BLANCO)
            ax1.tick_params(colors=TEXTO_GRIS)
        else:
            ax1.text(0.5, 0.5, "Sin datos de ingresos", ha='center', va='center', color=TEXTO_GRIS)
            ax1.set_title("Ingresos por Mes", color=TEXTO_BLANCO)

        ax2.set_facecolor(FONDO_TARJETA)
        if distribucion:
            monedas = [d['moneda'] for d in distribucion]
            montos = [float(d['total_usd']) for d in distribucion]
            lista_colores = []
            for m in monedas:
                if m == "USD":
                    lista_colores.append(COLOR_ACENTO)
                elif m == "COP":
                    lista_colores.append(COLOR_AZUL)
                else:
                    lista_colores.append(COLOR_VERDE)

            ax2.pie(montos, labels=monedas, autopct='%1.1f%%', colors=lista_colores, textprops={'color': TEXTO_BLANCO})
            ax2.set_title("Cobros por Moneda", color=TEXTO_BLANCO)
        else:
            ax2.text(0.5, 0.5, "Sin datos de monedas", ha='center', va='center', color=TEXTO_GRIS)
            ax2.set_title("Distribución por Moneda", color=TEXTO_BLANCO)

        canvas = FigureCanvasTkAgg(fig, master=self.marco_graficas_container)
        canvas.draw()
        canvas.get_tk_widget().pack(fill="both", expand=True, padx=5, pady=5)

        def cerrar_grafica():
            plt.close(fig)

        self.marco_graficas_container.bind("<Destroy>", lambda e: cerrar_grafica())

    def on_orden_seleccionada(self, event):
        seleccion = self.arbol_ordenes.selection()
        if not seleccion:
            return
        elemento = self.arbol_ordenes.item(seleccion)
        id_orden = elemento['values'][0]
        self.abrir_detalle_orden(id_orden)

    def abrir_detalle_orden(self, id_orden):
        try:
            datos = self.bd.obtener_detalle_orden_pagos(id_orden)
            if not datos or datos.get('id') is None:
                messagebox.showerror("Error", f"No se encontró la orden #{id_orden}")
                return

            pagos = self.bd.listar_pagos_por_orden(id_orden)

            ventana = ctk.CTkToplevel(self.padre)
            ventana.title(f"Detalle de Orden #{id_orden}")
            ventana.geometry("700x500")
            ventana.resizable(False, False)

            def al_cerrar():
                ventana.destroy()

            ventana.protocol("WM_DELETE_WINDOW", al_cerrar)

            marco = ctk.CTkFrame(ventana, fg_color=FONDO_TARJETA)
            marco.pack(fill="both", expand=True, padx=20, pady=20)

            ctk.CTkLabel(marco, text=f"Orden #{id_orden}", font=("Inter", 16, "bold"), text_color=TEXTO_BLANCO).pack(anchor="w", pady=(0, 5))
            ctk.CTkLabel(marco, text=f"Cliente: {datos.get('cliente_nombre', 'N/A')}", text_color=TEXTO_GRIS).pack(anchor="w")
            ctk.CTkLabel(marco, text=f"Vehículo: {datos.get('vehiculo', 'N/A')}", text_color=TEXTO_GRIS).pack(anchor="w")
            ctk.CTkLabel(marco, text=f"Descripción: {datos.get('descripcion', '')[:60]}...", text_color=TEXTO_GRIS).pack(anchor="w", pady=(0, 10))

            marco_resumen = ctk.CTkFrame(marco, fg_color=FONDO_SIDEBAR, corner_radius=10)
            marco_resumen.pack(fill="x", pady=10)

            total = datos.get('total_orden_usd', 0) or 0
            pagado = datos.get('total_pagado', 0) or 0
            saldo = max(0, total - pagado)

            ctk.CTkLabel(marco_resumen, text=f"Total: ${total:.2f} USD", text_color=TEXTO_BLANCO).pack(side="left", padx=15, pady=5)
            ctk.CTkLabel(marco_resumen, text=f"Pagado: ${pagado:.2f} USD", text_color=COLOR_VERDE).pack(side="left", padx=15, pady=5)
            ctk.CTkLabel(marco_resumen, text=f"Saldo: ${saldo:.2f} USD", text_color=COLOR_ACENTO if saldo > 0 else COLOR_VERDE, font=("Inter", 12, "bold")).pack(side="left", padx=15, pady=5)

            if pagos:
                ctk.CTkLabel(marco, text="Historial de Pagos:", font=("Inter", 12, "bold"), text_color=TEXTO_BLANCO).pack(anchor="w", padx=10, pady=5)

                arbol_pagos = ttk.Treeview(
                    marco,
                    columns=("Monto", "Moneda", "Tasa", "Monto USD", "Fecha", "Método", "Referencia"),
                    show="headings"
                )
                encabezados = [("Monto", 100), ("Moneda", 80), ("Tasa", 80), ("Monto USD", 100), ("Fecha", 120), ("Método", 100), ("Referencia", 100)]
                for columna, width in encabezados:
                    arbol_pagos.heading(columna, text=columna)
                    arbol_pagos.column(columna, width=width, anchor="center")

                scroll = ttk.Scrollbar(marco, orient="vertical", command=arbol_pagos.yview)
                arbol_pagos.configure(yscrollcommand=scroll.set)
                arbol_pagos.pack(fill="both", expand=True, pady=5)
                scroll.pack(side="right", fill="y")

                for p in pagos:
                    arbol_pagos.insert("", "end", values=(
                        f"{p['monto_original']:.2f}",
                        p['moneda'],
                        f"{p['tasa_cambio']:.2f}",
                        f"{p['monto_ref_usd']:.2f}",
                        formatear_fecha(p["fecha_pago"]),
                        p['metodo_pago'],
                        p['referencia'] or ""
                    ))
            else:
                ctk.CTkLabel(marco, text="No hay pagos registrados para esta orden", text_color=TEXTO_GRIS).pack(pady=10)

            boton_cerrar = ctk.CTkButton(marco, text="Cerrar", fg_color=COLOR_ACENTO, text_color=TEXTO_BLANCO, command=ventana.destroy)
            boton_cerrar.pack(pady=10)

        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al cargar el detalle:\n{str(e)}")

    LINEA = "=" * 78
    SEPARADOR = "-" * 78

    def _generar_texto_reporte(self, filas, anio, mes, nombre_mes):
        """Arma el texto del recibo de recaudos, en columnas de ancho fijo."""
        total_cobrado = sum(float(f['cobrado_mes']) for f in filas)
        lineas = [
            self.LINEA,
            self.LINEA,
            "                 TALLER DON JULIO - REPORTE DE RECAUDOS".center(78),
            self.LINEA,
            f" Periodo:  {nombre_mes} {anio}",
            f" Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            f" Usuario:  {self.usuario_actual}",
            self.SEPARADOR,
            " DETALLE POR ORDEN",
            self.SEPARADOR,
            f" {'#':<3}{'CLIENTE':<20}{'PLACA':<11}{'TOTAL ORDEN':>13} {'COBRADO MES':>14} {'SALDO':>11} ",
            self.SEPARADOR,
        ]

        for i, f in enumerate(filas, start=1):
            total_orden = float(f['total_orden_usd'] or 0)
            cobrado_mes = float(f['cobrado_mes'] or 0)
            cobrado_total = float(f['cobrado_total'] or 0)
            saldo = total_orden - cobrado_total
            lineas.append(
                f" {i:<3}{str(f['cliente'])[:18]:<20}{str(f['placa']):<11}"
                f"${total_orden:>12,.2f} ${cobrado_mes:>13,.2f} ${saldo:>10,.2f} "
            )

        lineas += [
            self.SEPARADOR,
            " RESUMEN DEL MES",
            self.SEPARADOR,
            f" Ordenes con cobros:   {len(filas)}",
            f" Total cobrado (USD):  ${total_cobrado:>12,.2f}",
            self.LINEA,
            " NOTA: este reporte muestra el dinero cobrado (ingresos) del periodo.",
            " No descuenta el costo de los repuestos, la mano de obra ni los gastos",
            " generales, por lo que no equivale a la ganancia final del taller.",
            self.LINEA,
        ]
        return "\n".join(lineas), total_cobrado

    def abrir_reporte_recaudos(self):
        meses = self.bd.listar_meses_con_pagos()
        if not meses:
            messagebox.showinfo("Sin datos", "Todavia no hay pagos registrados.", parent=self.marco)
            return

        ventana = ctk.CTkToplevel(self.padre)
        ventana.title("Reporte de Recaudos")
        ventana.geometry("860x700")
        ventana.resizable(False, False)

        marco = ctk.CTkFrame(ventana, fg_color=FONDO_TARJETA)
        marco.pack(fill="both", expand=True, padx=15, pady=15)

        ctk.CTkLabel(marco, text="\U0001F4B5 Reporte de Recaudos por Mes", font=("Inter", 16, "bold"),
                     text_color=TEXTO_BLANCO).pack(anchor="w")
        ctk.CTkLabel(marco, text="Muestra el dinero cobrado en el mes elegido. No incluye costos ni "
                                 "gastos, por lo que no es la ganancia final del taller.",
                      font=("Inter", 11), text_color=TEXTO_GRIS, wraplength=790,
                      justify="left").pack(anchor="w", pady=(2, 10))

        barra = ctk.CTkFrame(marco, fg_color="transparent")
        barra.pack(fill="x")

        ctk.CTkLabel(barra, text="Mes:", text_color=TEXTO_BLANCO).pack(side="left", padx=(0, 6))
        etiquetas = [f"{MESES_ES[m['mes'] - 1]} {m['anio']}" for m in meses]
        self.mes_reporte_var = ctk.StringVar(value=etiquetas[0])
        self.lista_mes_reporte = ctk.CTkComboBox(barra, values=etiquetas, width=190, state="readonly",
                                                 variable=self.mes_reporte_var)
        self.lista_mes_reporte.pack(side="left")

        boton_guardar = ctk.CTkButton(barra, text="\U0001F4BE Guardar como Excel", fg_color=COLOR_VERDE,
                                      text_color=TEXTO_BLANCO, state="disabled",
                                      command=self._guardar_reporte_excel)
        boton_guardar.pack(side="left", padx=10)

        self.caja_reporte = ctk.CTkTextbox(marco, font=("Courier New", 12), wrap="none")
        self.caja_reporte.pack(fill="both", expand=True, pady=12)
        self.caja_reporte.configure(state="disabled")

        ctk.CTkButton(marco, text="Cerrar", fg_color=COLOR_ACENTO, text_color=TEXTO_BLANCO,
                      command=ventana.destroy).pack(pady=(0, 5))

        self._boton_guardar = boton_guardar
        self._meses_reporte = dict(zip(etiquetas, meses))
        self.mes_reporte_var.trace_add("write", self._mostrar_vista_previa)
        self._mostrar_vista_previa()

    def _mostrar_vista_previa(self, *_eventos):
        """Pinta el reporte del mes elegido; se refresca solo al cambiar de mes."""
        datos = self._meses_reporte.get(self.mes_reporte_var.get())
        if not datos:
            return

        filas = self.bd.obtener_recaudos_del_mes(int(datos['anio']), int(datos['mes']))
        if filas:
            texto, _ = self._generar_texto_reporte(filas, datos['anio'], datos['mes'],
                                                  MESES_ES[datos['mes'] - 1])
        else:
            texto = "No hay cobros registrados en el mes seleccionado."

        self.caja_reporte.configure(state="normal")
        self.caja_reporte.delete("1.0", ctk.END)
        self.caja_reporte.insert("1.0", texto)
        self.caja_reporte.configure(state="disabled")

        self._boton_guardar.configure(
            state="normal" if filas else "disabled",
            text="\U0001F4BE Guardar como Excel" if filas else "Sin datos para guardar"
        )

    def _generar_excel_reporte(self, filas, anio, mes, nombre_mes):
        """Arma el libro de Excel con el detalle por orden y el resumen del mes."""
        libro = Workbook()
        hoja = libro.active
        hoja.title = f"Recaudos {nombre_mes} {anio}"

        azul = PatternFill("solid", fgColor="1F4E78")
        gris = PatternFill("solid", fgColor="DDEBF7")
        blanco = Font(bold=True, color="FFFFFF")
        negrita = Font(bold=True)
        centrado = Alignment(horizontal="center", vertical="center")
        derecha = Alignment(horizontal="right", vertical="center")
        borde = Border(*(Side(style="thin", color="B0B0B0"),) * 4)
        formato_dinero = '#,##0.00'

        hoja.merge_cells("A1:F1")
        titulo = hoja["A1"]
        titulo.value = "TALLER DON JULIO - REPORTE DE RECAUDOS"
        titulo.font = Font(bold=True, size=14, color="FFFFFF")
        titulo.fill = azul
        titulo.alignment = centrado
        hoja.row_dimensions[1].height = 26

        hoja.merge_cells("A2:F2")
        hoja["A2"] = f"Periodo: {nombre_mes} {anio}   |   Generado: {datetime.now():%Y-%m-%d %H:%M}   |   Usuario: {self.usuario_actual}"
        hoja["A2"].alignment = centrado
        hoja["A2"].font = Font(italic=True)

        encabezados = ["#", "Cliente", "Placa", "Total Orden (USD)", "Cobrado Mes (USD)", "Saldo (USD)"]
        fila_encabezado = 4
        for columna, texto in enumerate(encabezados, start=1):
            celda = hoja.cell(row=fila_encabezado, column=columna, value=texto)
            celda.font = blanco
            celda.fill = azul
            celda.alignment = centrado
            celda.border = borde

        fila_actual = fila_encabezado + 1
        for i, f in enumerate(filas, start=1):
            total_orden = float(f['total_orden_usd'] or 0)
            cobrado_mes = float(f['cobrado_mes'] or 0)
            saldo = total_orden - float(f['cobrado_total'] or 0)
            valores = [i, f['cliente'], f['placa'], total_orden, cobrado_mes, saldo]
            for columna, valor in enumerate(valores, start=1):
                celda = hoja.cell(row=fila_actual, column=columna, value=valor)
                celda.border = borde
                if columna in (4, 5, 6):
                    celda.number_format = formato_dinero
                    celda.alignment = derecha
                elif columna in (1, 3):
                    celda.alignment = centrado
            fila_actual += 1

        fila_total = fila_actual
        hoja.cell(row=fila_total, column=1, value="TOTALES").font = negrita
        hoja.merge_cells(start_row=fila_total, start_column=1, end_row=fila_total, end_column=3)
        total_cobrado = sum(float(f['cobrado_mes'] or 0) for f in filas)
        total_saldos = sum(float(f['total_orden_usd'] or 0) - float(f['cobrado_total'] or 0) for f in filas)
        for columna, valor in ((4, sum(float(f['total_orden_usd'] or 0) for f in filas)),
                               (5, total_cobrado), (6, total_saldos)):
            celda = hoja.cell(row=fila_total, column=columna, value=valor)
            celda.font = negrita
            celda.number_format = formato_dinero
            celda.alignment = derecha
        for columna in range(1, 7):
            hoja.cell(row=fila_total, column=columna).fill = gris
            hoja.cell(row=fila_total, column=columna).border = borde

        fila_resumen = fila_total + 2
        hoja.cell(row=fila_resumen, column=1, value="RESUMEN DEL MES").font = negrita
        resumen = [
            ("Órdenes con cobros", len(filas)),
            ("Total cobrado (USD)", total_cobrado),
            ("Saldo pendiente (USD)", total_saldos),
        ]
        for offset, (etiqueta, valor) in enumerate(resumen, start=1):
            celda_etiqueta = hoja.cell(row=fila_resumen + offset, column=1, value=etiqueta)
            celda_etiqueta.font = negrita
            celda_valor = hoja.cell(row=fila_resumen + offset, column=2, value=valor)
            if isinstance(valor, float):
                celda_valor.number_format = formato_dinero
                celda_valor.alignment = derecha

        fila_nota = fila_resumen + len(resumen) + 2
        nota = hoja.cell(row=fila_nota, column=1,
                         value="Nota: este reporte muestra el dinero cobrado (ingresos) del periodo. "
                               "No descuenta el costo de los repuestos, la mano de obra ni los gastos "
                               "generales, por lo que no equivale a la ganancia final del taller.")
        nota.alignment = Alignment(wrap_text=True, vertical="top")
        hoja.merge_cells(start_row=fila_nota, start_column=1, end_row=fila_nota + 2, end_column=6)

        anchos = (5, 26, 12, 18, 19, 15)
        for columna, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        hoja.freeze_panes = "A5"

        return libro, total_cobrado

    def _guardar_reporte_excel(self):
        datos = self._meses_reporte.get(self.mes_reporte_var.get())
        if not datos:
            return

        filas = self.bd.obtener_recaudos_del_mes(int(datos['anio']), int(datos['mes']))
        if not filas:
            messagebox.showinfo("Sin datos", "No hay cobros registrados en el mes seleccionado.",
                                parent=self.marco)
            return

        nombre_mes = MESES_ES[datos['mes'] - 1]
        libro, total = self._generar_excel_reporte(filas, datos['anio'], datos['mes'], nombre_mes)
        sugerencia = f"reporte_recaudos_{nombre_mes.lower()}_{datos['anio']}.xlsx"
        ruta = filedialog.asksaveasfilename(
            parent=self.marco,
            title="Guardar reporte de recaudos",
            defaultextension=".xlsx",
            initialfile=sugerencia,
            filetypes=[("Libro de Excel", "*.xlsx"), ("Todos los archivos", "*.*")]
        )
        if not ruta:
            return

        try:
            libro.save(ruta)
        except OSError as e:
            messagebox.showerror("Error", f"No se pudo guardar el archivo:\n{e}", parent=self.marco)
            return

        self.bd.registrar_log(self.usuario_id, self.usuario_actual, "presupuestos", 0,
                              "CREATE", f"Reporte de recaudos generado: {os.path.basename(ruta)}")
        messagebox.showinfo("\u00c9xito",
                            f"Reporte guardado en:\n{ruta}\n\nTotal cobrado: ${total:,.2f} "
                            f"en {len(filas)} \u00f3rdenes.", parent=self.marco)
